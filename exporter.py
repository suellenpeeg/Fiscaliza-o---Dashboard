from pptx import Presentation
from pptx.util import Inches
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from io import BytesIO
import tempfile

def gerar_ppt(df, fig1, fig2):

    prs = Presentation()

    # Slide 1 - Capa
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Painel de Controle – Fiscalização 2026"
    subtitle = slide.placeholders[1]
    subtitle.text = "Relatório Executivo Gerado Automaticamente"

    # Slide 2 - KPIs
    slide_layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(slide_layout)

    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4))
    tf = txBox.text_frame
    tf.text = f"""
    Total de Ações: {df["TOTAL_ACOES"].sum()}
    Total Banco de Dados: {df["TOTAL_BD"].sum()}
    Índice de Conformidade: {df["CONFORMIDADE"].mean():.1f}%
    """

    # Slide 3 - Gráfico 1
    with tempfile.NamedTemporaryFile(suffix=".png") as tmp:
        fig1.write_image(tmp.name)
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.add_picture(tmp.name, Inches(1), Inches(1), width=Inches(8))

    # Slide 4 - Gráfico 2
    with tempfile.NamedTemporaryFile(suffix=".png") as tmp:
        fig2.write_image(tmp.name)
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.add_picture(tmp.name, Inches(1), Inches(1), width=Inches(8))

    output = BytesIO()
    prs.save(output)
    output.seek(0)

    return output
